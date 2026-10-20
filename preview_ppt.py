#!/usr/bin/env python3
"""把 PPTX 的形状坐标/文字/颜色还原成 HTML，用于在无 PowerPoint 的环境下目检版式。

只还原几何与样式，不是渲染引擎 —— 字体度量与 PowerPoint 仍有细微差异，
但足以发现重叠、错位、配色错误、图片变形这类问题。
"""
import base64, html, sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

EMU_IN = 914400
ALIGN = {PP_ALIGN.LEFT: "left", PP_ALIGN.CENTER: "center",
         PP_ALIGN.RIGHT: "right", None: "left"}
ANCH = {MSO_ANCHOR.TOP: "flex-start", MSO_ANCHOR.MIDDLE: "center",
        MSO_ANCHOR.BOTTOM: "flex-end", None: "flex-start"}


def hexc(fmt, default=None):
    try:
        if fmt.type is not None and fmt.fore_color.type is not None:
            return "#%02x%02x%02x" % fmt.fore_color.rgb[0:3] if False else \
                   "#" + str(fmt.fore_color.rgb)
    except Exception:
        pass
    return default


def render(path, out, scale=0.78):
    prs = Presentation(path)
    SW, SH = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    parts = ["""<!DOCTYPE html><html><head><meta charset="utf-8">
<title>PPTX 版式预览</title><style>
body{margin:0;background:#2a2a28;font-family:"Noto Sans SC",system-ui,sans-serif;padding:24px}
.wrap{max-width:%dpx;margin:0 auto}
.lbl{color:#999;font-size:12px;margin:18px 0 6px}
.slide{position:relative;background:#fcfcfb;box-shadow:0 4px 18px rgba(0,0,0,.4);
  overflow:hidden}
.sh{position:absolute;box-sizing:border-box}
.tx{position:absolute;display:flex;flex-direction:column;box-sizing:border-box;
  white-space:pre-wrap;word-break:break-word;overflow:visible}
img{position:absolute}
</style></head><body><div class="wrap">""" % int(SW * 96 * scale)]

    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f'<div class="lbl">第 {i} 页</div>')
        parts.append(f'<div class="slide" style="width:{SW*96*scale:.1f}px;'
                     f'height:{SH*96*scale:.1f}px">')
        for sh in slide.shapes:
            if sh.left is None:
                continue
            L = sh.left / EMU_IN * 96 * scale
            T = sh.top / EMU_IN * 96 * scale
            Wd = (sh.width or 0) / EMU_IN * 96 * scale
            Ht = (sh.height or 0) / EMU_IN * 96 * scale

            if sh.shape_type == 13 or sh.__class__.__name__ == "Picture":
                b64 = base64.b64encode(sh.image.blob).decode()
                parts.append(f'<img src="data:{sh.image.content_type};base64,{b64}" '
                             f'style="left:{L:.1f}px;top:{T:.1f}px;'
                             f'width:{Wd:.1f}px;height:{Ht:.1f}px">')
                continue

            if sh.has_table:
                tb = sh.table
                yy = T
                for r_i, row in enumerate(tb.rows):
                    rh = row.height / EMU_IN * 96 * scale
                    xx = L
                    for c_i, col in enumerate(tb.columns):
                        cwid = col.width / EMU_IN * 96 * scale
                        cell = tb.cell(r_i, c_i)
                        bgc = hexc(cell.fill, "#fff")
                        txt = cell.text
                        p0 = cell.text_frame.paragraphs[0]
                        al = ALIGN.get(p0.alignment, "left")
                        fs = next((r.font.size.pt for r in p0.runs if r.font.size), 12)
                        bold = any(r.font.bold for r in p0.runs)
                        col_ = next((("#" + str(r.font.color.rgb))
                                     for r in p0.runs
                                     if r.font.color and r.font.color.type is not None), "#0b0b0b")
                        parts.append(
                            f'<div class="tx" style="left:{xx:.1f}px;top:{yy:.1f}px;'
                            f'width:{cwid:.1f}px;height:{rh:.1f}px;background:{bgc};'
                            f'justify-content:center;align-items:'
                            f'{"center" if al=="center" else "flex-start"};'
                            f'text-align:{al};font-size:{fs*scale:.1f}px;color:{col_};'
                            f'font-weight:{"600" if bold else "400"};'
                            f'border-bottom:1px solid #e1e0d9;padding:0 6px">'
                            f'{html.escape(txt)}</div>')
                        xx += cwid
                    yy += rh
                continue

            # 普通形状
            style = (f'left:{L:.1f}px;top:{T:.1f}px;width:{Wd:.1f}px;height:{Ht:.1f}px;')
            fillc = hexc(sh.fill)
            if fillc:
                style += f"background:{fillc};"
            try:
                if sh.line.fill.type is not None and sh.line.color.type is not None:
                    style += f"border:1px solid #{sh.line.color.rgb};"
            except Exception:
                pass
            nm = sh.shape_type
            if nm is not None and "OVAL" in str(nm):
                style += "border-radius:50%;"
            elif nm is not None and "ROUNDED" in str(nm):
                style += "border-radius:9px;"
            elif nm is not None and "TRIANGLE" in str(nm):
                style += ("background:transparent;border-left:%dpx solid transparent;"
                          "border-right:%dpx solid transparent;border-top:%dpx solid %s;"
                          "height:0;width:0;" % (Wd/2, Wd/2, Ht, fillc or "#c3c2b7"))
            parts.append(f'<div class="sh" style="{style}"></div>')

            if sh.has_text_frame and sh.text_frame.text.strip():
                tf = sh.text_frame
                anc = ANCH.get(tf.vertical_anchor, "flex-start")
                inner = []
                for p in tf.paragraphs:
                    t = "".join(r.text for r in p.runs)
                    if not t:
                        continue
                    fs = next((r.font.size.pt for r in p.runs if r.font.size), 18)
                    bold = any(r.font.bold for r in p.runs)
                    c = next((("#" + str(r.font.color.rgb)) for r in p.runs
                              if r.font.color and r.font.color.type is not None), "#0b0b0b")
                    al = ALIGN.get(p.alignment, "left")
                    sa = p.space_after.pt * scale if p.space_after else 0
                    ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
                    inner.append(
                        f'<div style="font-size:{fs*scale:.1f}px;color:{c};'
                        f'font-weight:{"600" if bold else "400"};text-align:{al};'
                        f'margin-bottom:{sa:.1f}px;line-height:{ls}">'
                        f'{html.escape(t)}</div>')
                parts.append(
                    f'<div class="tx" style="left:{L:.1f}px;top:{T:.1f}px;'
                    f'width:{Wd:.1f}px;height:{Ht:.1f}px;justify-content:{anc}">'
                    + "".join(inner) + "</div>")
        parts.append("</div>")
    parts.append("</div></body></html>")
    open(out, "w", encoding="utf-8").write("".join(parts))
    print(f"已生成 {out}（{len(prs.slides._sldIdLst)} 页）")


if __name__ == "__main__":
    render(sys.argv[1] if len(sys.argv) > 1 else "电解液温度依赖性研究_汇报.pptx",
           sys.argv[2] if len(sys.argv) > 2 else "ppt_preview.html")
